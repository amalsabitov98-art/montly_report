"""Extract July and August booking rows from the authoritative workbook."""
import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).parent
WORKBOOK = Path(r"C:\Users\User\Downloads\Отчеты (3).xlsx")
COLUMNS = [
    "id", "tourists", "partner", "direction", "sales", "net",
    "profit", "received", "debt", "paid", "manager",
]
COLUMN_NUMBERS = [1, 4, 5, 8, 9, 10, 11, 12, 13, 14, 17]


def extract(sheet):
    rows = []
    for cells in sheet.iter_rows(min_row=2, max_col=max(COLUMN_NUMBERS)):
        booking_id = cells[0].value
        if booking_id in (None, ""):
            continue
        row = [cells[column - 1].value for column in COLUMN_NUMBERS]
        rows.append(row)
    return rows


def main():
    workbook = load_workbook(WORKBOOK, data_only=True, read_only=True)
    for sheet_name, filename in (("июль", "july-data.json"), ("август", "august-data.json")):
        payload = {
            "source": WORKBOOK.name,
            "sheet": sheet_name,
            "columns": COLUMNS,
            "rows": extract(workbook[sheet_name]),
            "excludedBookingIds": [],
        }
        if sheet_name == "август":
            payload["marketing"] = 1277
            payload["managerMarketingDeduction"] = 159.6
        (ROOT / filename).write_text(
            json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        print(filename, len(payload["rows"]))


if __name__ == "__main__":
    main()
